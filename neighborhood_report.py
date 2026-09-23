from io import BytesIO
from datetime import date
import math
import pandas as pd
import requests
import streamlit as st
from docx import Document
from docx.shared import Inches

CENSUS_REPORTER_API = 'https://api.censusreporter.org/1.0/data/show/latest'
NOMINATIM_URL = 'https://nominatim.openstreetmap.org/search'
OVERPASS_URL = 'https://overpass-api.de/api/interpreter'
WIKIMEDIA_API = 'https://commons.wikimedia.org/w/api.php'
OPENCRIME_CITY_INDEX_URL = 'https://www.opencrime.us/data/city-index.json'
USER_AGENT = 'RentFinderNeighborhoodReport/1.0'


def _float(v):
    try:
        x = float(v)
        return None if math.isnan(x) else x
    except Exception:
        return None


def _money(v):
    x = _float(v)
    return 'N/A' if x is None else f'${x:,.0f}'


def _pct(v):
    x = _float(v)
    return 'N/A' if x is None else f'{x:.1f}%'


@st.cache_data(ttl=86400, show_spinner=False)
def geocode_zip(zip_code):
    r = requests.get(
        NOMINATIM_URL,
        params={'postalcode': zip_code, 'country': 'United States', 'format': 'jsonv2', 'limit': 1},
        headers={'User-Agent': USER_AGENT},
        timeout=20,
    )
    r.raise_for_status()
    data = r.json()
    if not data:
        return {'lat': None, 'lon': None, 'label': f'ZIP {zip_code}'}
    return {
        'lat': _float(data[0].get('lat')),
        'lon': _float(data[0].get('lon')),
        'label': data[0].get('display_name', f'ZIP {zip_code}'),
    }


@st.cache_data(ttl=86400, show_spinner=False)
def census_profile(zip_code):
    geoid = f'86000US{zip_code}'
    r = requests.get(
        CENSUS_REPORTER_API,
        params={'table_ids': 'B25064,B19013,B25018,B15003,B08301', 'geo_ids': geoid},
        headers={'User-Agent': USER_AGENT},
        timeout=30,
    )
    r.raise_for_status()
    est = r.json().get('data', {}).get(geoid, {}).get('estimate', {})

    def val(t, k):
        try:
            return est[t][k]
        except Exception:
            return None

    edu_total = _float(val('B15003', 'B15003001'))
    bachelors = sum((_float(val('B15003', k)) or 0) for k in ('B15003022','B15003023','B15003024','B15003025'))
    commute_total = _float(val('B08301', 'B08301001'))

    def share(k):
        x = _float(val('B08301', k))
        return None if not commute_total or x is None else x / commute_total * 100

    return {
        'median_gross_rent': _float(val('B25064', 'B25064001')),
        'median_household_income': _float(val('B19013', 'B19013001')),
        'median_rooms': _float(val('B25018', 'B25018001')),
        'bachelors_plus_pct': bachelors / edu_total * 100 if edu_total else None,
        'drive_alone_pct': share('B08301003'),
        'public_transit_pct': share('B08301010'),
        'walk_pct': share('B08301019'),
        'work_from_home_pct': share('B08301021'),
    }


@st.cache_data(ttl=21600, show_spinner=False)
def nearby_osm(lat, lon, radius_m):
    if lat is None or lon is None:
        return {'transit': [], 'schools': [], 'apartments': []}

    query = f'''[out:json][timeout:25];(
      nwr(around:{radius_m},{lat},{lon})["public_transport"];
      nwr(around:{radius_m},{lat},{lon})["railway"~"station|halt|tram_stop|subway_entrance"];
      nwr(around:{radius_m},{lat},{lon})["amenity"~"school|college|university"];
      nwr(around:{radius_m},{lat},{lon})["building"="apartments"];
    );out center tags 120;'''

    r = requests.post(OVERPASS_URL, data={'data': query}, headers={'User-Agent': USER_AGENT}, timeout=35)
    r.raise_for_status()
    out = {'transit': [], 'schools': [], 'apartments': []}

    for item in r.json().get('elements', []):
        tags = item.get('tags', {})
        center = item.get('center', {})
        entry = {
            'name': tags.get('name') or tags.get('official_name') or 'Unnamed',
            'lat': _float(item.get('lat') or center.get('lat')),
            'lon': _float(item.get('lon') or center.get('lon')),
        }
        if 'public_transport' in tags or tags.get('railway') in ('station','halt','tram_stop','subway_entrance'):
            e = dict(entry); e['type'] = tags.get('public_transport') or tags.get('railway') or 'transit'; out['transit'].append(e)
        if tags.get('amenity') in ('school','college','university'):
            e = dict(entry); e['type'] = tags.get('amenity'); out['schools'].append(e)
        if tags.get('building') == 'apartments':
            e = dict(entry); e['type'] = 'apartment building'; out['apartments'].append(e)

    for key in out:
        seen, clean = set(), []
        for x in out[key]:
            sig = (x['name'].lower(), round(x['lat'] or 0, 5), round(x['lon'] or 0, 5))
            if sig in seen: continue
            seen.add(sig); clean.append(x)
            if len(clean) >= 12: break
        out[key] = clean
    return out


@st.cache_data(ttl=43200, show_spinner=False)
def nearby_photos(lat, lon, radius_m):
    if lat is None or lon is None:
        return []
    r = requests.get(
        WIKIMEDIA_API,
        params={
            'action':'query','format':'json','generator':'geosearch','ggsprimary':'all','ggsnamespace':6,
            'ggsradius':min(radius_m,10000),'ggscoord':f'{lat}|{lon}','ggslimit':30,
            'prop':'imageinfo','iiprop':'url','iiurlwidth':900,'origin':'*'
        },
        headers={'User-Agent': USER_AGENT}, timeout=25,
    )
    r.raise_for_status()
    photos = []
    preferred = ('apartment','building','residential','housing','street','residence')
    for page in r.json().get('query', {}).get('pages', {}).values():
        info = (page.get('imageinfo') or [{}])[0]
        url = info.get('thumburl') or info.get('url')
        if not url: continue
        title = page.get('title','').replace('File:','')
        photos.append({'title': title, 'url': url, 'score': int(any(w in title.lower() for w in preferred))})
    photos.sort(key=lambda x: x['score'], reverse=True)
    return photos[:6]


@st.cache_data(ttl=86400, show_spinner=False)
def crime_index():
    r = requests.get(OPENCRIME_CITY_INDEX_URL, headers={'User-Agent': USER_AGENT}, timeout=45)
    r.raise_for_status()
    payload = r.json()
    if isinstance(payload, list): return payload
    if isinstance(payload, dict):
        for k in ('data','cities','records'):
            if isinstance(payload.get(k), list): return payload[k]
    return []


def find_crime(city, state):
    city, state = str(city or '').lower().strip(), str(state or '').upper().strip()
    try:
        records = crime_index()
    except Exception:
        return None
    for rec in records:
        rc = str(rec.get('name') or rec.get('city') or '').lower().strip()
        rs = str(rec.get('stateAbbr') or rec.get('state_abbr') or rec.get('state') or '').upper().strip()
        if rc == city and rs == state:
            return rec
    return None


def _image_bytes(url):
    try:
        r = requests.get(url, headers={'User-Agent': USER_AGENT}, timeout=20)
        r.raise_for_status(); return BytesIO(r.content)
    except Exception:
        return None


def build_docx(label, rent, census, nearby, crime, photos):
    doc = Document()
    doc.add_heading('Neighborhood Rental Report', 0)
    doc.add_paragraph(label)
    doc.add_paragraph(f"Generated: {date.today().strftime('%B %d, %Y')}")

    def table(rows):
        t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
        t.rows[0].cells[0].text='Metric'; t.rows[0].cells[1].text='Value'
        for a,b in rows:
            c=t.add_row().cells; c[0].text=str(a); c[1].text=str(b)

    doc.add_heading('Rental Market', level=1)
    table([
        ('Typical asking rent (ZORI)', _money(rent['latest_rent'])),
        ('Observation month', rent['latest_date'].strftime('%B %Y')),
        ('Census median gross rent', _money(census.get('median_gross_rent'))),
    ])

    doc.add_heading('Housing Size', level=1)
    table([('Median rooms', census.get('median_rooms') or 'N/A'), ('Exact square footage', 'Not available from the free area-level sources used here')])
    doc.add_paragraph('Median rooms is shown as a housing-size proxy. Exact square footage requires listing/property-level data.')

    doc.add_heading('Transportation', level=1)
    table([
        ('Public transit', _pct(census.get('public_transit_pct'))), ('Walk', _pct(census.get('walk_pct'))),
        ('Drive alone', _pct(census.get('drive_alone_pct'))), ('Work from home', _pct(census.get('work_from_home_pct'))),
    ])
    for x in nearby['transit'][:10]: doc.add_paragraph(f"{x['name']} — {x['type']}", style='List Bullet')

    doc.add_heading('Crime / Safety Context', level=1)
    if crime:
        table([
            ('Violent crime rate', f"{crime.get('violentCrimeRate','N/A')} per 100,000"),
            ('Property crime rate', f"{crime.get('propertyCrimeRate','N/A')} per 100,000"),
            ('Murder rate', f"{crime.get('murderRate','N/A')} per 100,000"),
        ])
        doc.add_paragraph('City-level FBI-derived figures are context only and do not predict risk at a specific address.')
    else:
        doc.add_paragraph('No matching city-level crime record was available.')

    doc.add_heading('Education', level=1)
    table([
        ("Adults 25+ with bachelor's degree or higher", _pct(census.get('bachelors_plus_pct'))),
        ('Median household income', _money(census.get('median_household_income'))),
        ('Nearby schools / colleges found', len(nearby['schools'])),
    ])
    for x in nearby['schools'][:10]: doc.add_paragraph(f"{x['name']} — {x['type']}", style='List Bullet')

    doc.add_heading('Nearby Apartment Buildings', level=1)
    if nearby['apartments']:
        for x in nearby['apartments'][:10]: doc.add_paragraph(x['name'], style='List Bullet')
    else:
        doc.add_paragraph('No named apartment buildings were returned nearby.')

    doc.add_heading('Nearby Public Photos', level=1)
    doc.add_paragraph('These are nearby Wikimedia Commons building/neighborhood photos, not guaranteed to be active listing photos.')
    for p in photos[:4]:
        b = _image_bytes(p['url'])
        if b:
            try:
                doc.add_picture(b, width=Inches(5.2)); doc.add_paragraph(p['title'])
            except Exception: pass

    doc.add_heading('Sources & Limitations', level=1)
    doc.add_paragraph('Rent: Zillow Research ZORI. Demographics/housing/transportation/education: Census Reporter/ACS. Nearby places: OpenStreetMap. Crime: OpenCrime city data processed from FBI Crime Data Explorer. Photos: Wikimedia Commons.')
    doc.add_paragraph('Datasets use different dates and geographic boundaries. This report is informational, not a property appraisal, school-quality score, or prediction of safety.')

    out=BytesIO(); doc.save(out); return out.getvalue()


def render_neighborhood_report(resolve_location, load_zillow_rent_data, lookup_zip, date_columns):
    st.markdown('''
    <div class="hero">
        <h1>📊 Neighborhood Report</h1>
        <p>Rent, housing size, transportation, crime context, education, nearby apartments and public photos.</p>
    </div>
    ''', unsafe_allow_html=True)

    with st.form('neighborhood_report_form'):
        location = st.text_input('Address or ZIP code', placeholder='07102 or 123 Main St, Newark, NJ')
        radius = st.select_slider('Nearby search radius', options=[1,2,3,5], value=3, format_func=lambda x:f'{x} km')
        go = st.form_submit_button('Generate neighborhood report', type='primary', use_container_width=True)

    if not go: return
    if not location.strip():
        st.error('Enter an address or ZIP code.'); return

    try:
        with st.spinner('Building report from public data sources...'):
            zip_code, resolved_label = resolve_location(location)
            df = load_zillow_rent_data(); row = lookup_zip(df, zip_code); cols = [c for c in date_columns(df) if pd.notna(row.get(c))]
            latest_col = cols[-1]; latest = float(row[latest_col]); latest_date = pd.to_datetime(latest_col)
            recent = cols[-24:]
            history = pd.DataFrame({'Month':pd.to_datetime(recent),'Typical Asking Rent':[float(row[c]) for c in recent]}).set_index('Month')
            rent = {'latest_rent':latest,'latest_date':latest_date,'city':str(row.get('City','') or ''),'state':str(row.get('State','') or ''),'history':history}

            census = census_profile(zip_code)
            geo = geocode_zip(zip_code)
            nearby = nearby_osm(geo['lat'], geo['lon'], radius*1000)
            photos = nearby_photos(geo['lat'], geo['lon'], radius*1000)
            crime = find_crime(rent['city'], rent['state'])

        st.success(f'Report generated for ZIP {zip_code}')

        st.subheader('Rental market')
        c1,c2,c3 = st.columns(3)
        c1.metric('Typical asking rent', _money(latest))
        c2.metric('Census median gross rent', _money(census.get('median_gross_rent')))
        c3.metric('Median household income', _money(census.get('median_household_income')))
        st.line_chart(history, use_container_width=True)

        st.subheader('Housing size')
        a,b = st.columns(2)
        a.metric('Median rooms', census.get('median_rooms') or 'N/A')
        b.metric('Exact square footage', 'Not available')
        st.caption('Median rooms is an area-level housing-size proxy. Exact unit square footage requires listing/property-level data.')

        st.subheader('Transportation')
        a,b,c,d = st.columns(4)
        a.metric('Public transit', _pct(census.get('public_transit_pct')))
        b.metric('Walk', _pct(census.get('walk_pct')))
        c.metric('Drive alone', _pct(census.get('drive_alone_pct')))
        d.metric('Work from home', _pct(census.get('work_from_home_pct')))
        if nearby['transit']:
            tdf=pd.DataFrame(nearby['transit'])[['name','type']]; tdf.columns=['Nearby transit','Type']; st.dataframe(tdf, hide_index=True, use_container_width=True)

        st.subheader('Crime / safety context')
        if crime:
            a,b,c = st.columns(3)
            a.metric('Violent crime', f"{crime.get('violentCrimeRate','N/A')} / 100k")
            b.metric('Property crime', f"{crime.get('propertyCrimeRate','N/A')} / 100k")
            c.metric('Murder', f"{crime.get('murderRate','N/A')} / 100k")
            st.caption('City-level FBI-derived statistics; not a prediction of risk at a specific property.')
        else:
            st.info('No matching city-level crime record was available.')

        st.subheader('Education')
        a,b = st.columns(2)
        a.metric("Bachelor's degree or higher", _pct(census.get('bachelors_plus_pct')))
        b.metric('Nearby schools / colleges', len(nearby['schools']))
        if nearby['schools']:
            edf=pd.DataFrame(nearby['schools'])[['name','type']]; edf.columns=['Nearby education','Type']; st.dataframe(edf, hide_index=True, use_container_width=True)

        st.subheader('Nearby apartment buildings')
        if nearby['apartments']:
            adf=pd.DataFrame(nearby['apartments'])[['name','type']]; adf.columns=['Apartment building','Type']; st.dataframe(adf, hide_index=True, use_container_width=True)
        else:
            st.caption('No named apartment-building features were returned nearby.')

        st.subheader('Nearby public photos')
        st.caption('Public neighborhood/building images from Wikimedia Commons. These are not guaranteed to be photos of currently available apartments.')
        if photos:
            cols2=st.columns(2)
            for i,p in enumerate(photos):
                with cols2[i%2]: st.image(p['url'], caption=p['title'], use_container_width=True)
        else:
            st.info('No suitable public photos were found nearby.')

        report=build_docx(resolved_label, rent, census, nearby, crime, photos)
        st.download_button('Download full neighborhood report (.docx)', report, file_name=f'neighborhood_report_{zip_code}_{date.today().isoformat()}.docx', mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document', type='primary', use_container_width=True)

    except Exception as exc:
        st.error(f'Could not generate the report: {exc}')
