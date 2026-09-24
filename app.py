from io import BytesIO
import re
from datetime import date, datetime
from pathlib import Path
from copy import deepcopy

import pandas as pd
import requests
import streamlit as st
from pypdf import PdfReader
from docx import Document
from neighborhood_report import render_neighborhood_report

# -------------------------------------------------
# Public rental data
# -------------------------------------------------

ZORI_ZIP_URL = (
    "https://files.zillowstatic.com/research/public_csvs/zori/"
    "Zip_zori_uc_sfrcondomfr_sm_month.csv"
)

CENSUS_GEOCODER_URL = (
    "https://geocoding.geo.census.gov/geocoder/geographies/onelineaddress"
)

st.set_page_config(
    page_title="RentFinder",
    page_icon="🏠",
    layout="centered",
)

st.markdown(
    """
    <style>
        .block-container {
            max-width: 950px;
            padding-top: 2rem;
            padding-bottom: 3rem;
        }

        .hero {
            padding: 1.8rem;
            border-radius: 18px;
            background: linear-gradient(135deg, #111827, #1f2937);
            color: white;
            margin-bottom: 1.5rem;
        }

        .hero h1 {
            margin: 0;
            font-size: 2.5rem;
        }

        .hero p {
            margin-top: .55rem;
            margin-bottom: 0;
            color: #d1d5db;
            font-size: 1.05rem;
        }

        .rent-card, .lease-card {
            border: 1px solid #e5e7eb;
            border-radius: 18px;
            padding: 1.4rem;
            margin-top: 1rem;
        }

        .rent-big {
            font-size: 2.8rem;
            font-weight: 800;
            margin: .15rem 0;
        }

        .muted {
            color: #6b7280;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# RENT FINDER FUNCTIONS
# =========================================================

def extract_zip(text: str):
    match = re.search(r"\b(\d{5})(?:-\d{4})?\b", text)
    return match.group(1) if match else None


@st.cache_data(ttl=60 * 60 * 12, show_spinner=False)
def load_zillow_rent_data():
    response = requests.get(ZORI_ZIP_URL, timeout=45)
    response.raise_for_status()

    df = pd.read_csv(
        BytesIO(response.content),
        dtype={"RegionName": str},
        low_memory=False,
    )

    df["RegionName"] = (
        df["RegionName"]
        .astype(str)
        .str.replace(r"\.0$", "", regex=True)
        .str.zfill(5)
    )

    return df


@st.cache_data(ttl=60 * 60 * 24, show_spinner=False)
def geocode_address_to_zip(address: str):
    params = {
        "address": address,
        "benchmark": "Public_AR_Current",
        "vintage": "Current_Current",
        "format": "json",
    }

    response = requests.get(
        CENSUS_GEOCODER_URL,
        params=params,
        timeout=20,
    )
    response.raise_for_status()

    matches = response.json().get("result", {}).get("addressMatches", [])

    if not matches:
        raise LookupError(
            "I couldn't match that address. Try entering the 5-digit ZIP code directly."
        )

    match = matches[0]
    matched_address = match.get("matchedAddress", "")
    zip_code = extract_zip(matched_address)

    if not zip_code:
        geographies = match.get("geographies", {})

        for geography_group in geographies.values():
            if not isinstance(geography_group, list):
                continue

            for record in geography_group:
                for key in ("ZIP", "ZCTA5", "ZCTA5CE20", "ZCTA5CE10"):
                    value = record.get(key)
                    if value and re.fullmatch(r"\d{5}", str(value)):
                        zip_code = str(value)
                        break

                if zip_code:
                    break

            if zip_code:
                break

    if not zip_code:
        raise LookupError(
            "The address was matched, but I couldn't determine its ZIP code. "
            "Enter the ZIP code directly."
        )

    return zip_code, matched_address or address


def resolve_location(location: str):
    direct_zip = extract_zip(location)

    if direct_zip and len(location.strip()) <= 10:
        return direct_zip, f"ZIP {direct_zip}"

    if direct_zip and ("," in location or any(c.isalpha() for c in location)):
        return direct_zip, location.strip()

    return geocode_address_to_zip(location.strip())


def date_columns(df: pd.DataFrame):
    columns = []

    for col in df.columns:
        try:
            pd.to_datetime(col, format="%Y-%m-%d")
            columns.append(col)
        except (ValueError, TypeError):
            continue

    return sorted(columns)


def lookup_zip(df: pd.DataFrame, zip_code: str):
    match = df[df["RegionName"] == zip_code]

    if match.empty:
        raise LookupError(
            f"Zillow's ZIP-level rent dataset does not currently contain ZIP {zip_code}."
        )

    return match.iloc[0]


def money(value):
    if pd.isna(value):
        return "N/A"
    return f"${float(value):,.0f}"


def pct(value):
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):+.1f}%"


# =========================================================
# LEASE UPDATER FUNCTIONS
# =========================================================

MONTHS = (
    "January|February|March|April|May|June|July|August|"
    "September|October|November|December|Jan|Feb|Mar|Apr|May|"
    "Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
)

DATE_PATTERNS = [
    # January 29, 2026 / March 01, 2026 / April 1st 2018
    rf"\b(?:{MONTHS})\.?\s+\d{{1,2}}(?:st|nd|rd|th)?[,]?\s+\d{{4}}\b",
    # 1 day of July 2020
    rf"\b\d{{1,2}}(?:st|nd|rd|th)?\s+day\s+of\s+(?:{MONTHS})\.?\s+\d{{4}}\b",
    # 3/1/2026, 07-01-2020
    r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b",
    # 2026-03-01
    r"\b\d{4}-\d{1,2}-\d{1,2}\b",
]

DATE_RE = re.compile("|".join(f"(?:{p})" for p in DATE_PATTERNS), re.IGNORECASE)

LEASE_CONTEXT_WORDS = (
    "lease", "term", "begin", "commence", "commencement", "start", "ending",
    "end", "terminate", "termination", "expiration", "expires", "dated",
    "effective", "renewal", "lease year", "signature", "signed", "date:"
)
PERSON_CONTEXT_WORDS = (
    "dob", "birth", "born", "family", "occupant", "resident", "age"
)


def _all_docx_paragraphs(doc):
    """Yield normal, table, header, and footer paragraphs."""
    seen = set()

    def emit_paragraph(p):
        key = id(p._p)
        if key not in seen:
            seen.add(key)
            return p
        return None

    for p in doc.paragraphs:
        item = emit_paragraph(p)
        if item is not None:
            yield item

    def walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    item = emit_paragraph(p)
                    if item is not None:
                        yield item
                for nested in cell.tables:
                    yield from walk_table(nested)

    for table in doc.tables:
        yield from walk_table(table)

    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                item = emit_paragraph(p)
                if item is not None:
                    yield item
            for table in container.tables:
                yield from walk_table(table)


def extract_lease_text(uploaded_file):
    suffix = Path(uploaded_file.name).suffix.lower()
    raw = uploaded_file.getvalue()

    if suffix == ".pdf":
        reader = PdfReader(BytesIO(raw))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)

    elif suffix == ".docx":
        doc = Document(BytesIO(raw))
        text = "\n".join(p.text for p in _all_docx_paragraphs(doc))

    elif suffix == ".txt":
        text = raw.decode("utf-8", errors="replace")

    else:
        raise ValueError("Unsupported file type.")

    if not text.strip():
        raise ValueError(
            "No readable text was found. Scanned/image-only PDFs are not supported "
            "in this version."
        )
    return text


def _context_for_match(text, match, radius=100):
    left = max(0, match.start() - radius)
    right = min(len(text), match.end() + radius)
    return " ".join(text[left:right].split())


def classify_date_context(context):
    c = context.lower()
    if any(word in c for word in PERSON_CONTEXT_WORDS):
        # Birth dates should never be selected automatically.
        return "Personal / DOB", False
    if any(word in c for word in LEASE_CONTEXT_WORDS):
        return "Lease-related", True
    return "Review needed", False


def find_dates_with_context(text):
    found = []
    seen = set()
    for match in DATE_RE.finditer(text):
        value = match.group(0)
        key = value.lower()
        if key in seen:
            continue
        seen.add(key)
        context = _context_for_match(text, match)
        category, recommended = classify_date_context(context)
        found.append({
            "value": value,
            "context": context,
            "category": category,
            "recommended": recommended,
        })
    return found


def _replace_text_preserve_runs(paragraph, replacements):
    """
    Replace text while keeping paragraph formatting as intact as possible.
    Handles dates split across Word runs by rebuilding only the affected
    paragraph's visible text into the first run.
    """
    original = paragraph.text
    updated = original
    for old, new in sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True):
        updated = re.sub(re.escape(old), new, updated, flags=re.IGNORECASE)

    if updated == original:
        return False

    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)
    return True


def replace_dates_in_docx(raw_bytes, replacements):
    doc = Document(BytesIO(raw_bytes))
    changed = 0
    for paragraph in _all_docx_paragraphs(doc):
        if _replace_text_preserve_runs(paragraph, replacements):
            changed += 1

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue(), changed


def replace_dates_in_text(text, replacements):
    result = text
    for old, new in sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True):
        result = re.sub(re.escape(old), new, result, flags=re.IGNORECASE)
    return result


def build_docx_from_text(text, source_name):
    doc = Document()
    doc.add_heading("Updated Lease Draft", level=1)
    p = doc.add_paragraph()
    p.add_run("Source file: ").bold = True
    p.add_run(source_name)
    p = doc.add_paragraph()
    p.add_run("Generated: ").bold = True
    p.add_run(date.today().strftime("%B %d, %Y"))
    warning = doc.add_paragraph()
    warning.add_run(
        "DRAFT FOR REVIEW — Review every changed date and all lease terms before signing."
    ).bold = True
    doc.add_paragraph("")
    for line in text.splitlines():
        doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


def format_replacement_date(chosen_date, original):
    """Keep the replacement close to the original date style."""
    if re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", original):
        return chosen_date.strftime("%Y-%m-%d")
    if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{2,4}", original):
        return f"{chosen_date.month}/{chosen_date.day}/{chosen_date.year}"
    if re.fullmatch(r"\d{1,2}-\d{1,2}-\d{2,4}", original):
        return f"{chosen_date.month:02d}-{chosen_date.day:02d}-{chosen_date.year}"
    if re.search(r"\bday\s+of\b", original, re.I):
        return f"{chosen_date.day} day of {chosen_date.strftime('%B %Y')}"
    if re.search(r"\d{1,2}(st|nd|rd|th)", original, re.I):
        d = chosen_date.day
        suffix = "th" if 10 <= d % 100 <= 20 else {1:"st",2:"nd",3:"rd"}.get(d % 10, "th")
        return f"{chosen_date.strftime('%B')} {d}{suffix} {chosen_date.year}"
    # Preserve zero-padded day if original used it.
    day_match = re.search(r"\b(\d{1,2})\b", original)
    padded = bool(day_match and len(day_match.group(1)) == 2 and day_match.group(1).startswith("0"))
    day = f"{chosen_date.day:02d}" if padded else str(chosen_date.day)
    return f"{chosen_date.strftime('%B')} {day}, {chosen_date.year}"




# =========================================================
# UI: RENT FINDER
# =========================================================

def render_rent_finder():
    st.markdown(
        """
        <div class="hero">
            <h1>🏠 RentFinder</h1>
            <p>
                Find the typical asking rent for a U.S. ZIP code or street address.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    with st.form("rent_search"):
        location = st.text_input(
            "Location",
            placeholder="Example: 07102 or 123 Main St, Newark, NJ",
        )

        submitted = st.form_submit_button(
            "Find rent",
            type="primary",
            use_container_width=True,
        )

    if not submitted:
        return

    if not location.strip():
        st.error("Enter a ZIP code or full street address.")
        return

    try:
        with st.spinner("Finding rental market data..."):
            zip_code, resolved_location = resolve_location(location)
            df = load_zillow_rent_data()
            row = lookup_zip(df, zip_code)
            monthly_columns = date_columns(df)

        available = [
            col for col in monthly_columns
            if col in row.index and not pd.isna(row[col])
        ]

        if not available:
            raise LookupError(
                f"No usable rent observations were found for ZIP {zip_code}."
            )

        latest_col = available[-1]
        latest_value = float(row[latest_col])
        latest_date = pd.to_datetime(latest_col)

        previous_value = float(row[available[-2]]) if len(available) >= 2 else None

        mom_change = None
        if previous_value:
            mom_change = ((latest_value / previous_value) - 1) * 100

        target_yoy_date = latest_date - pd.DateOffset(years=1)
        dates_map = {pd.to_datetime(col): col for col in available}

        yoy_value = None

        if dates_map:
            closest = min(
                dates_map.keys(),
                key=lambda d: abs((d - target_yoy_date).days),
            )

            if abs((closest - target_yoy_date).days) <= 45:
                yoy_value = float(row[dates_map[closest]])

        yoy_change = None
        if yoy_value:
            yoy_change = ((latest_value / yoy_value) - 1) * 100

        st.success(f"Rental data found for ZIP {zip_code}")

        st.markdown(
            f"""
            <div class="rent-card">
                <div class="muted">Typical monthly asking rent</div>
                <div class="rent-big">{money(latest_value)}</div>
                <div class="muted">
                    ZIP {zip_code} • {latest_date.strftime("%B %Y")}
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        c1, c2 = st.columns(2)
        c1.metric("Month-over-month", pct(mom_change))
        c2.metric("Year-over-year", pct(yoy_change))

        st.caption(f"Search resolved as: {resolved_location}")

        recent = available[-24:]
        history = pd.DataFrame(
            {
                "Month": pd.to_datetime(recent),
                "Typical Asking Rent": [float(row[col]) for col in recent],
            }
        ).set_index("Month")

        st.subheader("Rent trend")
        st.line_chart(history, use_container_width=True)

    except Exception as exc:
        st.error(str(exc))


# =========================================================
# UI: LEASE UPDATER
# =========================================================

def render_lease_updater():
    st.markdown(
        """
        <div class="hero">
            <h1>📄 Lease Updater</h1>
            <p>
                Upload a lease, review the dates found, and create a new draft
                without overwriting the original document.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.warning(
        "DRAFTING TOOL — Changing lease dates can change legal rights and obligations. "
        "The app keeps the original file untouched and requires you to review the "
        "dates before generating a new draft."
    )

    uploaded = st.file_uploader(
        "Upload lease",
        type=["docx", "pdf", "txt"],
        accept_multiple_files=False,
    )

    if not uploaded:
        st.info("Upload a DOCX, PDF, or TXT lease to begin. DOCX preserves the original layout best.")
        return

    try:
        raw = uploaded.getvalue()
        with st.spinner("Reading lease and detecting dates..."):
            text = extract_lease_text(uploaded)
            detected = find_dates_with_context(text)

        st.success(f"Loaded {uploaded.name} — found {len(detected)} unique date(s).")

        if Path(uploaded.name).suffix.lower() == ".pdf":
            st.info(
                "PDF text can be detected, but the generated Word draft cannot preserve the "
                "original PDF layout. Upload the original DOCX when available for best results."
            )

        with st.expander("Preview extracted lease text"):
            st.text_area("Lease text", text, height=300, disabled=True, label_visibility="collapsed")

        if not detected:
            st.info("No supported date formats were detected.")
            return

        st.subheader("1. Review detected dates")
        st.caption(
            "Lease-related dates are preselected. Personal/DOB dates are intentionally left unselected."
        )

        replacements = {}
        for i, item in enumerate(detected):
            with st.container(border=True):
                c1, c2 = st.columns([1.15, 1])
                with c1:
                    st.markdown(f"**{item['value']}**")
                    st.caption(f"{item['category']} · …{item['context']}…")
                    selected = st.checkbox(
                        "Update this date",
                        value=item["recommended"],
                        key=f"lease_select_{uploaded.name}_{i}",
                    )
                with c2:
                    chosen = st.date_input(
                        "New date",
                        value=date.today(),
                        key=f"lease_newdate_{uploaded.name}_{i}",
                        disabled=not selected,
                    )
                    if selected:
                        replacements[item["value"]] = format_replacement_date(chosen, item["value"])

        st.subheader("2. Confirm changes")
        if replacements:
            preview_rows = [{"Current date": old, "New date": new} for old, new in replacements.items()]
            st.dataframe(pd.DataFrame(preview_rows), use_container_width=True, hide_index=True)
        else:
            st.info("Select at least one lease date above.")

        if st.button(
            "Generate updated lease draft",
            type="primary",
            use_container_width=True,
            disabled=not bool(replacements),
        ):
            suffix = Path(uploaded.name).suffix.lower()

            if suffix == ".docx":
                output_bytes, changed_paragraphs = replace_dates_in_docx(raw, replacements)
                if changed_paragraphs == 0:
                    raise ValueError(
                        "The dates were detected, but Word stored them in a structure that could not "
                        "be safely replaced. Please send this document for inspection."
                    )
            else:
                updated_text = replace_dates_in_text(text, replacements)
                output_bytes = build_docx_from_text(updated_text, uploaded.name)
                changed_paragraphs = None

            st.session_state["updated_lease_docx"] = output_bytes
            st.session_state["updated_lease_name"] = (
                f"{Path(uploaded.name).stem}_UPDATED_DRAFT_{date.today().isoformat()}.docx"
            )
            st.session_state["updated_lease_changes"] = replacements.copy()

        if "updated_lease_docx" in st.session_state:
            st.success("Updated draft created. The original uploaded lease was not changed.")
            st.subheader("3. Download")
            st.write("Changes included in this draft:")
            st.dataframe(
                pd.DataFrame([
                    {"Original": k, "Replacement": v}
                    for k, v in st.session_state["updated_lease_changes"].items()
                ]),
                use_container_width=True,
                hide_index=True,
            )
            st.download_button(
                "Download updated lease draft (.docx)",
                data=st.session_state["updated_lease_docx"],
                file_name=st.session_state["updated_lease_name"],
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )

    except Exception as exc:
        st.error(str(exc))




# =========================================================
# MAIN NAVIGATION
# =========================================================

st.sidebar.title("RentFinder")

page = st.sidebar.radio(
    "Choose a tool",
    [
        "🏠 Rent Finder",
        "📄 Lease Updater",
        "📊 Neighborhood Report",
    ],
)

if page == "🏠 Rent Finder":
    render_rent_finder()
elif page == "📄 Lease Updater":
    render_lease_updater()
else:
    render_neighborhood_report(resolve_location, load_zillow_rent_data, lookup_zip, date_columns)

st.sidebar.caption(
    "Rent data: Zillow Research ZORI. "
    "Address lookup: U.S. Census Bureau."
)
